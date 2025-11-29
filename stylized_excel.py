import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def styling_excel(excel_file):
    book = load_workbook(excel_file)
    ws = book.worksheets[0]


    header_fill = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")  # Мягкий зелёный
    header_font = Font(size=11, bold=True, italic=True, color="FF1F497D")  # Тёмно-синий текст
    cell_font = Font(size=10, color="FF333333")  # Тёмно-серый текст для обычных ячеек


    thin_border = Border(
        left=Side(style='thin', color="FFAAAAAA"),
        right=Side(style='thin', color="FFAAAAAA"),
        top=Side(style='thin', color="FFAAAAAA"),
        bottom=Side(style='thin', color="FFAAAAAA")
    )


    for col in ["A", "B", "C", "D"]:
        cell = ws[f"{col}1"]
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border


    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.font = cell_font
            cell.alignment = Alignment(vertical="center", wrap_text=True,horizontal="center")
            cell.border = thin_border

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 25

    path_name = os.path.join('excel_file', "Объявление.xlsx")
    book.save(path_name)
    print("Excel файл успешно сохранен!")


